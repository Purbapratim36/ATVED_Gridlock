from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_workflow_doc():
    doc = Document()
    
    # Title
    title = doc.add_heading('ATVED Prototype - Complete Execution Workflow', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph("This document outlines the step-by-step technical execution flow of the Automated Traffic Violation Enforcement & Detection (ATVED) system, from camera input to the final E-Challan generation.\n")
    
    # Step 1
    doc.add_heading('Step 1: Frame Ingestion & Capture', level=1)
    p = doc.add_paragraph('The execution begins at the edge. The system connects to a standard CCTV camera or an IP webcam feed. It extracts individual image frames (typically at 30 frames per second) from the live video stream and pre-processes them for the AI engines.')
    
    # Step 2
    doc.add_heading('Step 2: The Triple-Model AI Inference Layer', level=1)
    doc.add_paragraph('The single image frame is simultaneously passed through three distinct YOLOv8 Deep Learning models running in parallel on the GPU:')
    doc.add_paragraph('• Primary Model: Scans the entire frame to detect and map bounding boxes for all vehicles (Cars, Motorcycles, Trucks, Buses).', style='List Bullet')
    doc.add_paragraph('• Helmet & Tripling Model: A specialized, fine-tuned model that specifically scans the scene for helmets, people riding without helmets, and motorcycles carrying more than two passengers.', style='List Bullet')
    doc.add_paragraph('• Seatbelt Model: Another specialized model trained to pierce through windshield glare to detect the presence or absence of seatbelts on drivers.', style='List Bullet')

    # Step 3
    doc.add_heading('Step 3: Temporal Tracking (ByteTrack)', level=1)
    doc.add_paragraph('Raw AI detections are essentially "dumb" frame-by-frame boxes. The system utilizes the ByteTrack tracking algorithm to assign a unique, persistent ID to every vehicle. This allows the system to follow a specific car as it moves across the screen over time, rather than treating it as a new car in every frame.')

    # Step 4
    doc.add_heading('Step 4: The Mathematical Rule Engine', level=1)
    doc.add_paragraph('Once the AI identifies the objects and the tracker links them over time, the data is passed to the core Rule Engine. This engine applies hard mathematical logic to determine if a violation occurred:')
    doc.add_paragraph('• Speed Detector: Calculates the pixel-displacement of a vehicle\'s centroid across consecutive frames, converting it to kilometres-per-hour to check if it exceeds the speed limit.', style='List Bullet')
    doc.add_paragraph('• Spatial Logic: Checks if a "Motorcycle" bounding box perfectly overlaps with a "No Helmet" bounding box or a "Triple Riding" bounding box.', style='List Bullet')
    doc.add_paragraph('• Red Light Detector: Monitors an invisible Y-coordinate "Stop Line" on the screen. If a vehicle\'s bounding box crosses this line while the system state is set to RED, it triggers an event.', style='List Bullet')

    # Step 5
    doc.add_heading('Step 5: Confidence Scoring & Filtration', level=1)
    doc.add_paragraph('To absolutely eliminate false positives, the Rule Engine calculates a final Confidence Score for the violation. If the AI was only 40% confident, or if the violation only lasted for a fraction of a second, the event is immediately discarded. Only violations with a confidence score exceeding 75% are permitted to pass through the filter.')

    # Step 6
    doc.add_heading('Step 6: Backend API Serialization', level=1)
    doc.add_paragraph('A confirmed violation is packaged into a JSON payload containing the vehicle\'s tracking ID, the violation type (e.g., TRIPLE_RIDING), the confidence score, and the exact timestamp. This payload is fired via an HTTP POST request to the ATVED Backend API Server (api_server.py).')

    # Step 7
    doc.add_heading('Step 7: Database & Dashboard Updates', level=1)
    doc.add_paragraph('The final step occurs in the backend infrastructure:')
    doc.add_paragraph('• Database Storage: The API server permanently logs the E-Challan into the PostgreSQL database, automatically looking up the mocked license plate to deduct points from the driver\'s digital score.', style='List Bullet')
    doc.add_paragraph('• Live Socket Transmission: The server pushes the new data into a Redis cache stream.', style='List Bullet')
    doc.add_paragraph('• Real-Time Interface: The Police Web Dashboard, listening to the stream, instantly flashes red and updates the UI with the new ticket, while an automated SMS notification is fired to the offender.', style='List Bullet')

    doc.add_heading('Summary', level=2)
    doc.add_paragraph('This entire 7-step pipeline executes in a fraction of a second, completely autonomously, transforming a cheap street camera into an intelligent, unblinking traffic enforcement officer.')

    doc.save(r'C:\Users\mahan\OneDrive\Documents\Desktop\Projects\Gridlock_sunny\ATVED_Execution_Workflow.docx')
    print("Doc created successfully!")

if __name__ == "__main__":
    create_workflow_doc()
